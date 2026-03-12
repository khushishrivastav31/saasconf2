"""openpyxl workbook builder for analytics revenue exports.

This module is imported lazily and called via asyncio.to_thread() to avoid
blocking the event loop during workbook construction.
"""

import io
from decimal import Decimal

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from app.schemas.analytics import RevenueTrackingResponse


def build_revenue_workbook(revenue: RevenueTrackingResponse) -> bytes:
    """Build an .xlsx workbook with revenue data. Returns raw bytes."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    _sheet_revenue_by_option(wb, revenue)
    _sheet_summary(wb, revenue)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="2563EB")  # blue-600
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_BOLD = Font(bold=True)
_RIGHT = Alignment(horizontal="right")
_CURRENCY_FMT = "#,##0.00"


def _header_row(ws, headers: list[str]) -> None:
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _autofit(ws) -> None:  # type: ignore[no-untyped-def]
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)


_SUBROW_FILL = PatternFill("solid", fgColor="F3F4F6")  # gray-100
_OPTION_FILL = PatternFill("solid", fgColor="DBEAFE")  # blue-100


def _sheet_revenue_by_option(wb: openpyxl.Workbook, revenue: RevenueTrackingResponse) -> None:
    ws = wb.create_sheet("Revenue by Option")
    headers = [
        "Booking Option",
        "Category",
        "Period",
        "Period Start",
        "Period End",
        "Attendees",
        "Subtotal",
        "Tax",
        "Total",
        "Currency",
    ]
    _header_row(ws, headers)

    row = 2
    for item in revenue.by_booking_option:
        # Option summary row (bold, light blue background)
        opt_name_cell = ws.cell(row, 1, item.booking_option_name)
        opt_name_cell.font = _BOLD
        opt_name_cell.fill = _OPTION_FILL
        ws.cell(row, 2, item.category).fill = _OPTION_FILL
        ws.cell(row, 3, "All Periods").fill = _OPTION_FILL
        ws.cell(row, 4, "").fill = _OPTION_FILL
        ws.cell(row, 5, "").fill = _OPTION_FILL
        att_cell = ws.cell(row, 6, item.attendee_count)
        att_cell.alignment = _RIGHT
        att_cell.font = _BOLD
        att_cell.fill = _OPTION_FILL
        for col, val in [(7, item.subtotal_revenue), (8, item.tax_amount), (9, item.total_revenue)]:
            c = ws.cell(row, col, float(val))
            c.number_format = _CURRENCY_FMT
            c.alignment = _RIGHT
            c.font = _BOLD
            c.fill = _OPTION_FILL
        ws.cell(row, 10, item.currency).fill = _OPTION_FILL
        row += 1

        # Period sub-rows (indented, light gray)
        for period in item.periods:
            ws.cell(row, 1, f"  {period.period_name}").fill = _SUBROW_FILL
            ws.cell(row, 2, "").fill = _SUBROW_FILL
            ws.cell(row, 3, period.period_name).fill = _SUBROW_FILL
            ws.cell(row, 4, str(period.period_start) if period.period_start else "").fill = _SUBROW_FILL
            ws.cell(row, 5, str(period.period_end) if period.period_end else "").fill = _SUBROW_FILL
            pa_cell = ws.cell(row, 6, period.attendee_count)
            pa_cell.alignment = _RIGHT
            pa_cell.fill = _SUBROW_FILL
            for col, val in [
                (7, period.subtotal_revenue),
                (8, period.tax_amount),
                (9, period.total_revenue),
            ]:
                c = ws.cell(row, col, float(val))
                c.number_format = _CURRENCY_FMT
                c.alignment = _RIGHT
                c.fill = _SUBROW_FILL
            ws.cell(row, 10, item.currency).fill = _SUBROW_FILL
            row += 1

    # Totals row
    ws.cell(row, 1, "TOTAL").font = _BOLD
    ws.cell(row, 6, revenue.total_registrations).alignment = _RIGHT
    subtotal_sum = sum(item.subtotal_revenue for item in revenue.by_booking_option)
    tax_sum = sum(item.tax_amount for item in revenue.by_booking_option)
    s_cell = ws.cell(row, 7, float(subtotal_sum))
    s_cell.number_format = _CURRENCY_FMT
    s_cell.font = _BOLD
    s_cell.alignment = _RIGHT
    t_cell = ws.cell(row, 8, float(tax_sum))
    t_cell.number_format = _CURRENCY_FMT
    t_cell.font = _BOLD
    t_cell.alignment = _RIGHT
    tot_cell = ws.cell(row, 9, float(revenue.registration_revenue))
    tot_cell.number_format = _CURRENCY_FMT
    tot_cell.font = _BOLD
    tot_cell.alignment = _RIGHT
    ws.cell(row, 10, revenue.currency)

    _autofit(ws)


def _sheet_summary(wb: openpyxl.Workbook, revenue: RevenueTrackingResponse) -> None:
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18

    rows: list[tuple[str, Decimal | int | str]] = [
        ("Currency", revenue.currency),
        ("", ""),
        ("Registration Revenue", revenue.registration_revenue),
        ("Sponsor Revenue", revenue.sponsor_revenue),
        ("Total Revenue (paid)", revenue.total_revenue),
        ("Outstanding (receivables)", revenue.outstanding_revenue),
        ("", ""),
        ("Total Registrations", revenue.total_registrations),
        ("Paid Registrations", revenue.paid_registrations),
        ("Cancelled Registrations", revenue.cancelled_registrations),
        ("Sponsor Invoice Count", revenue.sponsor_invoice_count),
    ]

    ws.cell(1, 1, "Financial Summary").font = Font(bold=True, size=14)

    for i, (label, value) in enumerate(rows, start=3):
        ws.cell(i, 1, label)
        cell = ws.cell(i, 2, float(value) if isinstance(value, Decimal) else value)
        if isinstance(value, Decimal):
            cell.number_format = _CURRENCY_FMT
            cell.alignment = _RIGHT
